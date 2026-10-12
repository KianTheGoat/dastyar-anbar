import csv
import io
import re
import zipfile
from datetime import UTC, date, datetime, timedelta
from zoneinfo import ZoneInfo

import jdatetime
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session

from app.config import Settings
from app.db.models import (
    AggregateSale,
    Category,
    ImportBatch,
    ImportMapping,
    SaleLine,
    Variant,
    Workspace,
)
from app.security import audit, decrypt, encrypt, lock_workspace, utcnow

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10_000
FIELDS = {
    "inventory": ["sku", "name", "category", "color", "size", "available"],
    "sales": [
        "source",
        "order_id",
        "line_id",
        "sku",
        "sold_at",
        "quantity",
        "status",
        "returned_quantity",
    ],
    "aggregate": ["sku", "net_units", "period_start", "period_end"],
}
ALIASES = {
    "sku": ["کد کالا", "شناسه کالا"],
    "name": ["نام کالا", "نام محصول"],
    "category": ["دسته بندی", "دسته‌بندی"],
    "color": ["رنگ"],
    "size": ["سایز"],
    "available": ["موجودی", "موجودی قابل فروش"],
    "order_id": ["شماره سفارش"],
    "line_id": ["شماره ردیف"],
    "sold_at": ["تاریخ فروش"],
    "quantity": ["تعداد فروش"],
    "returned_quantity": ["تعداد مرجوعی"],
    "status": ["وضعیت"],
    "net_units": ["فروش خالص"],
    "period_start": ["شروع بازه"],
    "period_end": ["پایان بازه"],
}
STATUS_MAP = {
    v: v for v in ["preparing", "shipping", "completed", "pending", "cancelled", "refunded"]
}
STATUS_MAP.update(
    {
        "در حال آماده‌سازی": "preparing",
        "در حال آماده سازی": "preparing",
        "در حال ارسال": "shipping",
        "تکمیل شده": "completed",
        "تکمیل‌شده": "completed",
        "لغو شده": "cancelled",
        "لغوشده": "cancelled",
        "مسترد شده": "refunded",
        "processing": "preparing",
        "on-hold": "pending",
        "failed": "cancelled",
    }
)
DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def clean_headers(rows: list[list]) -> tuple[list[str], list[list]]:
    if not rows or not rows[0] or len(rows[0]) > 64:
        raise HTTPException(400, "invalid_headers")
    headers = [str(value or "").strip() for value in rows[0]]
    if any(not h or len(h) > 128 for h in headers) or len(set(headers)) != len(headers):
        raise HTTPException(400, "invalid_headers")
    data = [row for row in rows[1:] if any(value not in (None, "") for value in row)]
    if not data or len(data) > MAX_ROWS:
        raise HTTPException(400, "row_limit_or_empty")
    return headers, data


def read_file(raw: bytes, filename: str, sheet: str = "") -> tuple[list[str], list[list]]:
    if len(raw) > MAX_BYTES:
        raise HTTPException(413, "file_too_large")
    try:
        if filename.lower().endswith(".csv"):
            text = raw.decode("utf-8-sig")
            reader = csv.reader(io.StringIO(text))
            rows = []
            for row in reader:
                if len(rows) > MAX_ROWS or len(row) > 64:
                    raise HTTPException(400, "row_limit_or_empty")
                rows.append(row)
        elif filename.lower().endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                files = archive.infolist()
                if len(files) > 512 or sum(f.file_size for f in files) > 25 * 1024 * 1024:
                    raise HTTPException(400, "unsafe_workbook")
                for member in files:
                    name = member.filename.lower()
                    if (
                        name.startswith("/")
                        or ".." in name.split("/")
                        or "vbaproject" in name
                        or "externallinks" in name
                        or member.flag_bits & 1
                    ):
                        raise HTTPException(400, "unsafe_workbook")
                    if name.endswith(".xml"):
                        content = archive.read(member)
                        if b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper():
                            raise HTTPException(400, "unsafe_workbook")
            workbook = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            try:
                if len(workbook.sheetnames) > 10:
                    raise HTTPException(400, "unsafe_workbook")
                ws = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]
                if (ws.max_row and ws.max_row > MAX_ROWS + 1) or (
                    ws.max_column and ws.max_column > 64
                ):
                    raise HTTPException(400, "row_limit_or_empty")
                rows = []
                for row in ws.iter_rows(
                    max_row=ws.max_row or MAX_ROWS + 1, max_col=ws.max_column or 64
                ):
                    if any(cell.data_type == "f" for cell in row):
                        raise HTTPException(400, "formula_not_allowed")
                    rows.append([cell.value for cell in row])
            finally:
                workbook.close()
        else:
            raise HTTPException(400, "csv_or_xlsx_required")
        return clean_headers(rows)
    except (UnicodeError, csv.Error, zipfile.BadZipFile, ValueError, KeyError, OSError):
        raise HTTPException(400, "invalid_file") from None


def proposed_mapping(headers: list[str], kind: str, saved: dict | None = None) -> dict:
    result = {}
    for field in FIELDS[kind]:
        for label in [(saved or {}).get(field), field] + ALIASES.get(field, []):
            if label in headers:
                result[field] = label
                break
    return result


def as_count(value) -> int:
    if isinstance(value, bool) or value is None:
        raise ValueError("invalid_count")
    text = str(value).translate(DIGITS).strip()
    if re.fullmatch(r"[0-9]{1,3}(,[0-9]{3})+", text):
        text = text.replace(",", "")
    if not re.fullmatch(r"[0-9]+", text) or len(text) > 7 or int(text) > 1_000_000:
        raise ValueError("invalid_count")
    return int(text)


def as_text(value, maximum: int = 128, default: str | None = None) -> str:
    if value in (None, "") and default is not None:
        return default
    if not isinstance(value, str):
        # Numeric identifiers in Excel have already lost leading zeros.
        raise ValueError("text_required")
    value = value.strip()
    if not value or len(value) > maximum or any(ord(c) < 32 for c in value):
        raise ValueError("invalid_text")
    if value.startswith(("=", "+", "@")):
        raise ValueError("formula_not_allowed")
    return value


def as_datetime(value, calendar: str, timezone: str) -> datetime:
    zone = ZoneInfo(timezone)
    if isinstance(value, datetime):
        result = value
    elif isinstance(value, date):
        result = datetime.combine(value, datetime.min.time())
    else:
        text = str(value or "").translate(DIGITS).strip()
        if calendar == "jalali":
            match = re.fullmatch(
                r"(\d{4})[-/](\d{2})[-/](\d{2})(?:[ T](\d{2}):(\d{2})(?::(\d{2}))?)?", text
            )
            if not match:
                raise ValueError("date_format")
            y, m, d, h, minute, second = match.groups()
            day = jdatetime.date(int(y), int(m), int(d)).togregorian()
            result = datetime(
                day.year, day.month, day.day, int(h or 0), int(minute or 0), int(second or 0)
            )
        else:
            if not re.match(r"^\d{4}-\d{2}-\d{2}(?:T| |$)", text):
                raise ValueError("date_format")
            result = datetime.fromisoformat(text)
    if result.tzinfo is None:
        result = result.replace(tzinfo=zone)
    return result.astimezone(UTC)


def normalize(
    headers: list[str],
    rows: list[list],
    kind: str,
    mapping: dict,
    calendar: str = "gregorian",
    timezone: str = "Asia/Tehran",
    finalized: bool = False,
    status_mapping: dict | None = None,
) -> tuple[list[dict], list[dict]]:
    if (
        kind not in FIELDS
        or calendar not in {"gregorian", "jalali"}
        or timezone not in {"UTC", "Asia/Tehran"}
    ):
        raise HTTPException(400, "invalid_import_options")
    required = {
        "inventory": {"sku", "name", "category", "available"},
        "sales": {"order_id", "line_id", "sku", "sold_at", "quantity"},
        "aggregate": set(FIELDS["aggregate"]),
    }[kind]
    if kind == "sales" and not finalized:
        required.add("status")
    if any(not isinstance(k, str) or not isinstance(v, str) for k, v in mapping.items()):
        raise HTTPException(400, "invalid_mapping")
    if (
        set(mapping) - set(FIELDS[kind])
        or not required.issubset(mapping)
        or any(v not in headers for v in mapping.values())
        or len(set(mapping.values())) != len(mapping)
    ):
        raise HTTPException(400, "invalid_mapping")
    statuses = STATUS_MAP | (status_mapping or {})
    if len(statuses) > 100 or any(
        not isinstance(v, str) or v not in set(STATUS_MAP.values()) for v in statuses.values()
    ):
        raise HTTPException(400, "invalid_status_mapping")
    result, errors, seen = [], [], set()
    index = {field: headers.index(label) for field, label in mapping.items()}
    for number, row in enumerate(rows, 2):
        values = {field: row[i] if i < len(row) else None for field, i in index.items()}
        output = {}
        try:
            for field in FIELDS[kind]:
                value = values.get(field)
                if field in {"available", "quantity", "net_units", "returned_quantity"}:
                    output[field] = as_count(
                        value
                        if value is not None
                        else (0 if field == "returned_quantity" else None)
                    )
                elif field in {"sold_at", "period_start", "period_end"}:
                    converted = as_datetime(value, calendar, timezone)
                    output[field] = (
                        converted.isoformat()
                        if field == "sold_at"
                        else converted.astimezone(ZoneInfo(timezone)).date().isoformat()
                    )
                elif field == "status":
                    label = str(value or "").strip()
                    if value in (None, "") and finalized:
                        output[field] = "completed"
                    elif label in statuses:
                        output[field] = statuses[label]
                    else:
                        raise ValueError("unknown_status")
                else:
                    default = {"color": "بدون رنگ", "size": "بدون سایز", "source": "manual"}.get(
                        field
                    )
                    maximum = {
                        "name": 250,
                        "category": 120,
                        "color": 80,
                        "size": 80,
                        "source": 80,
                    }.get(field, 128)
                    if (
                        field == "size"
                        and isinstance(value, (int, float))
                        and not isinstance(value, bool)
                    ):
                        value = str(value)
                    output[field] = as_text(value, maximum, default)
            if kind == "sales":
                field = "returned_quantity"
                if output[field] > output["quantity"]:
                    raise ValueError("return_exceeds_sale")
                key = (output["source"], output["order_id"], output["line_id"])
            else:
                key = output["sku"]
            if key in seen:
                field = "sku" if kind != "sales" else "line_id"
                raise ValueError("duplicate_row")
            seen.add(key)
            result.append(output)
        except (ValueError, TypeError, OverflowError):
            errors.append({"row": number, "field": field, "code": "invalid_value"})
            if len(errors) >= 100:
                break
    if kind == "aggregate" and result:
        periods = {(v["period_start"], v["period_end"]) for v in result}
        if len(periods) != 1 or result[0]["period_start"] > result[0]["period_end"]:
            errors.append({"row": 0, "field": "period_start", "code": "inconsistent_period"})
    return result, errors


def save_preview(
    db: Session,
    workspace: Workspace,
    settings: Settings,
    kind: str,
    rows: list[dict],
    mapping: dict,
    snapshot_at: datetime | None = None,
    commit: bool = True,
) -> dict:
    locked = lock_workspace(db, workspace.id)
    if kind != "inventory":
        known = set(db.scalars(select(Variant.sku).where(Variant.workspace_id == workspace.id)))
        if any(v["sku"] not in known for v in rows):
            raise HTTPException(400, "unknown_sku")
    if kind == "inventory":
        existing = dict(
            db.execute(
                select(Variant.sku, Variant.inventory_updated_at).where(
                    Variant.workspace_id == workspace.id
                )
            ).all()
        )
        if len(set(existing) | {v["sku"] for v in rows}) > MAX_ROWS:
            raise HTTPException(400, "workspace_variant_limit")
        snapshot_at = snapshot_at or utcnow()
        if snapshot_at.tzinfo is None or snapshot_at > utcnow() + timedelta(minutes=5):
            raise HTTPException(400, "invalid_snapshot_time")
        if any(
            existing.get(v["sku"], datetime.min.replace(tzinfo=UTC)) > snapshot_at for v in rows
        ):
            raise HTTPException(409, "stale_inventory")
    # Workspace lock captures the version for apply; no raw file is kept.
    db.execute(
        delete(ImportBatch).where(
            ImportBatch.workspace_id == workspace.id,
            ImportBatch.kind == kind,
            ImportBatch.applied.is_(False),
        )
    )
    batch = ImportBatch(
        workspace_id=workspace.id,
        kind=kind,
        revision=locked.data_revision,
        payload=encrypt(
            settings,
            {
                "rows": rows,
                "mapping": mapping,
                "snapshot_at": snapshot_at.isoformat() if snapshot_at else None,
            },
        ),
        row_count=len(rows),
        expires_at=utcnow() + timedelta(hours=24),
    )
    db.add(batch)
    if commit:
        db.commit()
    else:
        db.flush()
    return {
        "id": str(batch.id),
        "kind": kind,
        "row_count": len(rows),
        "sample": rows[:5],
        "expires_at": batch.expires_at.isoformat(),
        "sales_mode": workspace.sales_mode,
        "changes_sales_mode": kind != "inventory"
        and workspace.sales_mode != ("aggregate" if kind == "aggregate" else "detail"),
    }


def apply_preview(
    db: Session, workspace_id, batch_id, settings: Settings, commit: bool = True
) -> dict:
    workspace = lock_workspace(db, workspace_id)
    batch = db.scalar(
        select(ImportBatch)
        .where(ImportBatch.id == batch_id, ImportBatch.workspace_id == workspace_id)
        .with_for_update()
    )
    if not batch:
        raise HTTPException(404, "not_found")
    if batch.applied:
        return {"ok": True, "already_applied": True, "row_count": batch.row_count}
    if batch.expires_at <= utcnow() or not batch.payload:
        raise HTTPException(410, "preview_expired")
    if batch.revision != workspace.data_revision:
        raise HTTPException(409, "preview_outdated")
    payload = decrypt(settings, batch.payload)
    kind = batch.kind
    if kind == "bundle":
        from app.db.models import SheetConnection

        connection = db.get(SheetConnection, workspace_id)
        if not connection:
            raise HTTPException(410, "sheet_disconnected")
        for child_kind, child_rows in [
            ("inventory", payload["inventory"]),
            (payload["sales_kind"], payload["sales"]),
        ]:
            child = ImportBatch(
                workspace_id=workspace_id,
                kind=child_kind,
                revision=workspace.data_revision,
                row_count=len(child_rows),
                applied=False,
                expires_at=batch.expires_at,
                payload=encrypt(
                    settings,
                    {
                        "rows": child_rows,
                        "mapping": payload["maps"][child_kind],
                        "snapshot_at": payload["snapshot_at"],
                    },
                ),
            )
            db.add(child)
            db.flush()
            apply_preview(db, workspace_id, child.id, settings, commit=False)
        batch.applied = True
        batch.payload = None
        connection.last_synced_at = utcnow()
        audit(db, workspace_id, "sheet_sync", batch.row_count)
        db.commit()
        return {"ok": True, "already_applied": False, "row_count": batch.row_count}
    rows = payload["rows"]
    if kind == "inventory":
        categories = {
            c.name: c
            for c in db.scalars(select(Category).where(Category.workspace_id == workspace_id))
        }
        for name in {row["category"] for row in rows}:
            if name not in categories:
                c = Category(workspace_id=workspace_id, name=name, low_max=3, high_min=5)
                db.add(c)
                db.flush()
                categories[name] = c
        when = datetime.fromisoformat(payload["snapshot_at"])
        for row in rows:
            values = {k: row[k] for k in ["sku", "name", "color", "size", "available"]}
            values |= {
                "workspace_id": workspace_id,
                "category_id": categories[row["category"]].id,
                "inventory_updated_at": when,
                "replenishment": "none",
            }
            statement = insert(Variant).values(**values)
            db.execute(
                statement.on_conflict_do_update(
                    constraint="uq_variant_workspace_sku",
                    set_={
                        k: values[k]
                        for k in values
                        if k not in {"workspace_id", "sku", "replenishment"}
                    },
                )
            )
    else:
        variants = {
            v.sku: v.id
            for v in db.scalars(select(Variant).where(Variant.workspace_id == workspace_id))
        }
        if kind == "aggregate":
            db.execute(delete(AggregateSale).where(AggregateSale.workspace_id == workspace_id))
            for row in rows:
                db.add(
                    AggregateSale(
                        workspace_id=workspace_id,
                        variant_id=variants[row["sku"]],
                        net_units=row["net_units"],
                    )
                )
            workspace.sales_mode = "aggregate"
            workspace.aggregate_start = date.fromisoformat(rows[0]["period_start"])
            workspace.aggregate_end = date.fromisoformat(rows[0]["period_end"])
        else:
            keys = set(
                db.execute(
                    select(
                        SaleLine.source, SaleLine.external_order_id, SaleLine.external_line_id
                    ).where(SaleLine.workspace_id == workspace_id)
                ).all()
            )
            if len(keys | {(r["source"], r["order_id"], r["line_id"]) for r in rows}) > 100_000:
                raise HTTPException(400, "workspace_sale_limit")
            for row in rows:
                values = {
                    "workspace_id": workspace_id,
                    "variant_id": variants[row["sku"]],
                    "source": row["source"],
                    "external_order_id": row["order_id"],
                    "external_line_id": row["line_id"],
                    "sold_at": datetime.fromisoformat(row["sold_at"]),
                    "quantity": row["quantity"],
                    "returned_quantity": row["returned_quantity"],
                    "status": row["status"],
                }
                db.execute(
                    insert(SaleLine)
                    .values(**values)
                    .on_conflict_do_update(
                        constraint="uq_sale_external_line",
                        set_={
                            k: values[k]
                            for k in values
                            if k
                            not in {
                                "workspace_id",
                                "source",
                                "external_order_id",
                                "external_line_id",
                            }
                        },
                    )
                )
            workspace.sales_mode = "detail"
            workspace.aggregate_start = workspace.aggregate_end = None
            db.execute(delete(AggregateSale).where(AggregateSale.workspace_id == workspace_id))
    saved = db.get(ImportMapping, (workspace_id, kind))
    if saved:
        saved.mapping = payload["mapping"]
    else:
        db.add(ImportMapping(workspace_id=workspace_id, kind=kind, mapping=payload["mapping"]))
    batch.applied = True
    batch.payload = None
    workspace.data_revision += 1
    audit(db, workspace_id, "import_" + kind, len(rows))
    if commit:
        db.commit()
    else:
        db.flush()
    return {"ok": True, "already_applied": False, "row_count": len(rows)}


def csv_cell(value) -> str:
    text = str(value)
    candidate = text.lstrip(" \t\r\n\ufeff\u200b")
    if candidate.startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text
